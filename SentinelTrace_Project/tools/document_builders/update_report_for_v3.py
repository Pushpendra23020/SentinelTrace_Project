"""Apply small, traceable v3 code-structure corrections to the final report."""
from pathlib import Path

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT = PROJECT_ROOT / "deliverables" / "02_Project_Report" / "SentinelTrace_Final_Report.docx"


def replace_paragraph_text(document: Document, old: str, new: str) -> None:
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new)
            return
    # The helper is intentionally safe to re-run during document QA.
    if any(new in paragraph.text for paragraph in document.paragraphs):
        return
    raise ValueError(f"Expected report text not found: {old[:40]}...")


def replace_module_table(document: Document) -> None:
    updated_rows = [
        ["run.py", "Starts the local application", "One-command launcher"],
        ["src/sentineltrace/analyzer.py", "Parses, detects, and correlates events", "JSON-ready analysis report"],
        ["src/sentineltrace/rules.py", "Stores documented rule metadata and weights", "Auditable rule catalogue"],
        ["src/sentineltrace/server.py", "Validates HTTP input and serves local routes", "Local dashboard API"],
        ["web/index.html and tests/", "Dashboard interface and automated checks", "Evaluation evidence"],
    ]
    for table in document.tables:
        if any("app.py" in cell.text for row in table.rows for cell in row.cells):
            for row, values in zip(table.rows[1:], updated_rows):
                for cell, value in zip(row.cells, values):
                    cell.text = value
            return
    if any("src/sentineltrace/analyzer.py" in cell.text for table in document.tables for row in table.rows for cell in row.cells):
        return
    raise ValueError("Expected source-code modules table was not found.")


def remove_redundant_evidence_scope_row(document: Document) -> None:
    """Avoid a one-line table fragment at the top of the next report page."""

    for table in document.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == "Evidence scope":
                table._tbl.remove(row._tr)
                return


document = Document(REPORT)
replace_paragraph_text(
    document,
    "For every finding, the rule catalogue provides a documented weight. The entity risk equals the sum of associated finding weights, capped at 100.",
    "For every source IP, the rule catalogue contributes the configured weight of each triggered rule category once; the total is capped at 100.",
)
replace_paragraph_text(
    document,
    "The implementation deliberately keeps the module boundary compact so an individual student can understand and defend all code.",
    "The implementation uses small, named modules so an individual student can explain the flow from parsing through response generation.",
)
replace_paragraph_text(
    document,
    "Correlation groups all findings by source IP. The risk score is the sum of the documented rule weights, capped at 100.",
    "Correlation groups all findings by source IP. The risk score adds the documented weight of each triggered rule category once and caps the result at 100.",
)
replace_module_table(document)
remove_redundant_evidence_scope_row(document)
document.save(REPORT)
