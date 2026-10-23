from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT_ROOT=Path(__file__).resolve().parents[2]
REFERENCE=Path('/home/pushpendra/Downloads/88e336ad-180a-4e34-8ab5-8f4075d34845.docx')
OUTPUT=PROJECT_ROOT/'deliverables'/'01_Synopsis'/'SentinelTrace_Synopsis_Template_Format.docx'
TITLE='SentinelTrace An Explainable Local Log Based Threat Triage and Alert Prioritization System'

def set_text(paragraph, text):
    if paragraph.runs:
        first=paragraph.runs[0]
        first.text=text
        for run in paragraph.runs[1:]: run.text=''
    else: paragraph.add_run(text)

def set_cell_text(cell, text):
    set_text(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]: set_text(paragraph, '')

def delete_from(document, start_text):
    start=None
    for paragraph in document.paragraphs:
        if start_text in paragraph.text:
            start=paragraph._p
            break
    if start is None: raise ValueError(f'Missing template marker: {start_text}')
    body=document._element.body
    deleting=False
    for element in list(body):
        if element is start: deleting=True
        if deleting and element.tag != qn('w:sectPr'):
            body.remove(element)

def apply_body_format(document):
    section=document.sections[-1]
    section.top_margin=Pt(70.87); section.bottom_margin=Pt(35.43)
    section.left_margin=Pt(70.87); section.right_margin=Pt(35.43)
    normal=document.styles['Normal']
    normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.5

def heading(document, text):
    p=document.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(9)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run(text); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(15)
    return p

def body(document, text, bold_prefix=None):
    p=document.add_paragraph(); p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_after=Pt(6); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(11); r=p.add_run(text[len(bold_prefix):]); r.font.name='Times New Roman'; r.font.size=Pt(11)
    else:
        r=p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(11)
    return p

def bullet(document, text):
    p=document.add_paragraph(style='Normal'); p.paragraph_format.left_indent=Pt(24); p.paragraph_format.first_line_indent=Pt(-12); p.paragraph_format.space_after=Pt(3)
    r=p.add_run('• '+text); r.font.name='Times New Roman'; r.font.size=Pt(11)
    return p

def add_table(document, headers, rows, widths=None):
    caption=body(document, 'Table: '+headers[0]+' details', 'Table: '); caption.paragraph_format.space_before=Pt(6)
    t=document.add_table(rows=1, cols=len(headers))
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        border=OxmlElement(f'w:{edge}'); border.set(qn('w:val'),'single'); border.set(qn('w:sz'),'6'); border.set(qn('w:color'),'000000'); borders.append(border)
    t._tbl.tblPr.append(borders)
    for cell, text in zip(t.rows[0].cells,headers):
        cell.text=text
        for run in cell.paragraphs[0].runs: run.bold=True; run.font.name='Times New Roman'; run.font.size=Pt(10)
    for row in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,row):
            cell.text=str(text)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.15
                for run in p.runs: run.font.name='Times New Roman'; run.font.size=Pt(10)
    return t

def page_break(document): document.add_page_break()

doc=Document(REFERENCE)
delete_from(doc,'FORMAT OF SYNOPSIS')
apply_body_format(doc)

# Cover slots
for p in doc.paragraphs:
    if 'PROJECT TITLE' in p.text: set_text(p,TITLE)
    elif 'Computer Science & Engineering' in p.text: set_text(p,'Computer Science and Engineering')
for cell in doc.tables[0].rows[0].cells:
    if 'Submitted by' in cell.text: set_cell_text(cell,'Submitted by\nStudent Name: ____________________\nRoll No.: ____________________\nSection: ____________________')
    elif 'Supervision' in cell.text: set_cell_text(cell,'Under the Supervision of\nGuide Name: ____________________\nAssistant Professor')

# Index slots
index=doc.tables[1]
topics=['Introduction','Statement of Problem and Objectives','Literature Survey','Expected Outcome and Scope of the Project','Tentative Work Plan','Software and Hardware Required','References']
pages=['1 to 2','3','4','5','6','7','8']
for i,(topic,page_no) in enumerate(zip(topics,pages),1):
    index.rows[i].cells[0].text=f'{i:02d}'; index.rows[i].cells[1].text=topic; index.rows[i].cells[2].text=page_no

# Section 1 and 2
heading(doc,'1. INTRODUCTION')
body(doc,'Cybersecurity monitoring begins with evidence. Web servers and operating systems continuously generate access and authentication logs, but a small laboratory may lack a simple method to inspect them, connect related events, and decide which activity needs attention first. SentinelTrace is a local defensive application designed for this need. It accepts authorised Apache style access logs and Linux SSH authentication logs, converts recognised lines into structured events, applies transparent rules, and displays a prioritised incident timeline for human review.')
body(doc,'The project belongs to the cybersecurity and security operations domain. Its core purpose is log based threat triage: reducing a raw stream of records into a smaller set of explainable observations. The project does not scan targets, exploit weaknesses, execute command text, transmit logs to an external service, or conclude that an alert alone proves compromise. It is deliberately bounded as a local decision support prototype suitable for an academic mini project.')
body(doc,'The proposed implementation uses Python 3 standard library components for the local HTTP server, parsing, correlation, hashing, and JSON responses. HTML, CSS, and JavaScript provide the browser dashboard. This technology selection keeps the system portable and makes every module practical to explain and modify during viva.')
body(doc,'Important terms used in the project include an event, which is a normalised log record; a finding, which is a rule triggered by evidence; correlation, which groups related findings by source IP address; severity, which indicates priority; and confidence, which cautions the analyst about the strength of an indicator. MITRE ATT and CK technique identifiers are shown with findings to provide a common vocabulary for observed behaviour.')
body(doc,'The architecture has four logical stages. First, parsers recognise the supported access and authentication formats. Second, an explainable rule catalogue identifies selected indicators such as repeated password failures or suspicious request paths. Third, a correlation layer calculates a transparent risk score for each source IP. Finally, the dashboard and JSON export present the evidence, rationale, recommendation, and known limitation for an analyst.')
page_break(doc)
heading(doc,'1. INTRODUCTION CONTINUED')
body(doc,'Security log analysis is often the earliest practical way to observe password guessing, reconnaissance, or abnormal web requests. However, a single line rarely provides sufficient context. SentinelTrace therefore gives precedence to relationships between events. For example, five failed SSH passwords from one IP create a high severity brute force finding. A successful login after that same pattern creates a separate, higher priority correlation finding because the analyst may need to validate account ownership and session activity.')
body(doc,'The rule based approach is intentional. Machine learning methods can be valuable when suitable labelled data, careful evaluation, and ongoing model governance are available. For an individual mini project, deterministic rules have clear advantages: they are reproducible, inspectable, and easy to modify. A student can demonstrate how one rule works, change its threshold, and explain the effect on risk scoring without hiding the decision behind an opaque model.')
add_table(doc,['Component','Role in SentinelTrace'],[['Parser','Recognises Apache style web requests and common SSH accepted or failed password lines.'],['Rule engine','Detects six bounded indicators and attaches severity, confidence, MITRE mapping, and response guidance.'],['Correlation layer','Groups findings by IP address and ranks entities using a capped weighted score.'],['Dashboard','Shows event count, findings, severity summary, entities, and chronological evidence.'],['JSON export','Creates a reproducible incident report for documentation and review.']])
page_break(doc)

heading(doc,'2. STATEMENT OF PROBLEM AND OBJECTIVES')
body(doc,'Raw security logs are heterogeneous and high volume. Manual review can be slow, especially when related web and authentication events occur in different formats. Small educational or laboratory environments may not have a production SIEM, while generic scripts frequently produce alerts without evidence, correlation, confidence, or a clear recommended response. The problem is to provide a local, understandable, and safe mechanism that highlights selected indicators in authorised logs without overstating what those indicators mean.')
body(doc,'The existing manual approach requires the analyst to read every line, remember source IP addresses, and connect repeated behaviour across sources. It can miss relationships, offers no consistent prioritisation, and is difficult to reproduce during a review. SentinelTrace addresses these limitations through deterministic processing and an analyst oriented interface.')
body(doc,'The objectives of the project are:')
for x in ['Parse supported Apache style access logs and Linux SSH authentication logs into normalised events.','Detect at least six indicators: SSH brute force, successful login after repeated failures, path traversal, SQL injection indicators, command injection indicators, and web scanning.','Attach evidence, severity, confidence, MITRE ATT and CK mapping, and a recommended analyst response to every finding.','Correlate findings by source IP address and calculate an auditable risk score capped at 100.','Display a prioritised dashboard and export a complete JSON incident report.','Validate expected behaviour using synthetic, safe, and repeatable log data.']: bullet(doc,x)
page_break(doc)

heading(doc,'3. LITERATURE SURVEY')
body(doc,'NIST Special Publication 800 92, Guide to Computer Security Log Management, establishes the importance of collection, analysis, retention, and protection of security log data. It supports the project decision to treat log records as operational evidence and to distinguish collection from meaningful review. SentinelTrace narrows this broad guidance to a local, educational triage workflow.')
body(doc,'The OWASP Logging Cheat Sheet explains that application logs should support detection, investigation, and monitoring while being designed securely. Its guidance informs the project emphasis on evidence, source context, and safe handling. SentinelTrace does not create a complete application logging policy; instead, it consumes selected log output and presents suspicious patterns for review.')
body(doc,'MITRE ATT and CK provides a shared knowledge base for adversary behaviours. The project maps password guessing to T1110.001, valid accounts to T1078, exploitation of public facing applications to T1190, command and scripting interpreter behaviour to T1059, and active scanning to T1595. These mappings help a reviewer discuss behaviour consistently but do not establish attribution or intent.')
body(doc,'Existing enterprise SIEM products offer broad ingestion, storage, alerting, and correlation. They are useful in production but can be expensive, complex, and difficult to reproduce in a compact individual project. In contrast, basic single rule scripts are often transparent but lack a unified user interface and correlation. SentinelTrace occupies the middle ground: a small, local, explainable prototype with a bounded threat model.')
add_table(doc,['Source','Relevance to the Project'],[['NIST SP 800 92','Security log management principles and operational context.'],['OWASP Logging Cheat Sheet','Application logging, monitoring, and security considerations.'],['MITRE ATT and CK','Common terminology for the selected behavioural indicators.'],['OWASP SQL Injection Prevention Cheat Sheet','Context for treating suspicious query content as an indicator, not proof.']])
page_break(doc)

heading(doc,'4. EXPECTED OUTCOME AND SCOPE OF THE PROJECT')
body(doc,'The expected outcome is a working local dashboard that accepts authorised text logs, parses recognised events, detects six selected rule categories, prioritises source IP addresses, and displays an explainable timeline. The dashboard will show the rule, evidence, severity, confidence, MITRE identifier, source line, and recommended response. It will also download a JSON incident report that preserves the result and the system limitations.')
body(doc,'The intended users are students, faculty reviewers, and junior analysts working with safe laboratory logs. Practical use includes explaining how a threat triage workflow moves from raw log text to a review queue. The included synthetic demo corpus allows repeatable demonstration without exposing personal data or attacking any system.')
body(doc,'The scope is limited to two documented log patterns and a small rule catalogue. It is not a production SIEM, an endpoint protection tool, a vulnerability scanner, or an automated containment system. A finding remains a prompt for analyst validation. This boundary is central to the project ethics and prevents the prototype from being misrepresented.')
body(doc,'Future extensions may include timestamp normalisation and sliding windows, encrypted case storage, role based authentication, structured JSON log ingestion, Sigma rule import, and carefully evaluated anomaly detection with a public labelled dataset. These extensions are not represented as completed features in the current prototype.')
page_break(doc)

heading(doc,'5. TENTATIVE WORK PLAN')
body(doc,'The project follows the required ten week cycle. Each week has a visible deliverable so that progress can be demonstrated, documented, and reviewed by the mentor.')
add_table(doc,['Week','Major Activity','Deliverable'],[[1,'Domain finalisation and title options','Feasibility notes and title alternatives'],[2,'Synopsis preparation','Synopsis and initial references'],[3,'Mentor review and approval','Approved title and revised scope'],[4,'Requirements and system design','Architecture, data model, rule catalogue'],[5,'Implementation phase one','Log parser and initial rules'],[6,'Midterm review','At least 30 percent working demonstration'],[7,'Implementation phase two','Correlation, risk scoring, explainability, export'],[8,'Integration and testing','Functional tests and corrections'],[9,'Report preparation','Draft report, screenshots, plagiarism review'],[10,'Final submission and viva','Source code, final report, presentation, viva']])
body(doc,'Weekly work will be recorded in the student logbook with the date, completed tasks, issues faced, and mentor remarks. Any change to the approved scope will be discussed before implementation so that the final system remains feasible within the timeline.')
page_break(doc)

heading(doc,'6. SOFTWARE AND HARDWARE REQUIRED')
body(doc,'The project is intentionally lightweight and can be developed on a standard student laptop. The application uses only local resources and does not require paid cloud infrastructure, a database server, or specialised hardware.')
add_table(doc,['Category','Requirement','Purpose'],[['Operating system','Windows 10 or 11, Linux, or macOS','Local development and browser access.'],['Programming language','Python 3.10 or later','Parser, HTTP server, rules, correlation, and JSON API.'],['Frontend','HTML, CSS, JavaScript','Local browser dashboard and report download.'],['Development tools','VS Code or equivalent, Git, modern browser','Code editing, version control, testing, and demo.'],['Processor','Dual core 2 GHz or higher','Sufficient for bounded local log analysis.'],['Memory','4 GB RAM minimum, 8 GB recommended','Development tools and browser execution.'],['Storage','500 MB free space','Source code, documentation, screenshots, and safe demo logs.'],['Network','Not required for analysis','The server uses loopback address 127.0.0.1 only.']])
body(doc,'No personal or proprietary dataset is required. The demonstration uses synthetic logs stored with the project. If real logs are used later, they must be authorised, sanitised, and handled according to institutional requirements.')
page_break(doc)

heading(doc,'7. REFERENCES')
for ref in ['[1] K. Kent and M. Souppaya, Guide to Computer Security Log Management, NIST Special Publication 800 92, National Institute of Standards and Technology, 2006.','[2] OWASP Foundation, Logging Cheat Sheet, OWASP Cheat Sheet Series, https://cheatsheetseries.owasp.org/cheatsheets/Logging_Cheat_Sheet.html, accessed September 2026.','[3] MITRE ATT and CK, Brute Force Password Guessing T1110.001, https://attack.mitre.org/techniques/T1110/001/, accessed September 2026.','[4] MITRE ATT and CK, Valid Accounts T1078, https://attack.mitre.org/techniques/T1078/, accessed September 2026.','[5] MITRE ATT and CK, Exploit Public Facing Application T1190, https://attack.mitre.org/techniques/T1190/, accessed September 2026.','[6] OWASP Foundation, SQL Injection Prevention Cheat Sheet, https://cheatsheetseries.owasp.org/cheatsheets/SQL_Injection_Prevention_Cheat_Sheet.html, accessed September 2026.','[7] MITRE ATT and CK, Command and Scripting Interpreter T1059, https://attack.mitre.org/techniques/T1059/, accessed September 2026.']:
    p=body(doc,ref); p.alignment=WD_ALIGN_PARAGRAPH.LEFT

OUTPUT.parent.mkdir(exist_ok=True)
doc.save(OUTPUT)
